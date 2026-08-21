from pptx import Presentation


def parse_pptx(file_path: str) -> str:
    """Extract text from a PowerPoint file."""
    text = []
    try:
        prs = Presentation(file_path)
        for slide in prs.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    text.append(shape.text.strip())
    except Exception as e:
        raise ValueError(f"Failed to parse PPTX: {e}")
    return "\n".join(text)
