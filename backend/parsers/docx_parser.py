from docx import Document


def parse_docx(file_path: str) -> str:
    """Extract text from a Word document."""
    text = []
    try:
        doc = Document(file_path)
        for para in doc.paragraphs:
            if para.text.strip():
                text.append(para.text.strip())
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        text.append(cell.text.strip())
    except Exception as e:
        raise ValueError(f"Failed to parse DOCX: {e}")
    return "\n".join(text)
